"""extract_crd.py  (v3 - supports old format and new consolidated upstream format)
Usage: python extract_crd.py <excel_path> [crd_list]
  excel_path  Path to the SAP Ariba CRD Excel tracker
  crd_list    Comma-separated IDs (e.g. CRD-01,S-1,C-2) or 'all' (default: all)
Outputs JSON array to stdout.

Format detection:
  New format: has consolidated upstream sheets (Sourcing_Customizations etc.) + 'Overview' sheet
  Old format: individual CRD-XX sheets only (with optional 'CRD Overview' sheet)

New format upstream consolidated IDs:
  Sourcing     -> S-1, S-2  ... (Sourcing_Customizations)
  Contracts    -> C-1, C-2  ... (Contracts_Customizations)
  SPM          -> SPM-1 ...     (SPM_Customizations)
  Forms/dForms -> F-1, F-2  ... (Forms_Customizations)
  Savings      -> SF-1, SF-2 .. (Savings_Customizations)
"""
import openpyxl
import sys
import re
import json


# ---------------------------------------------------------------------------
# Adaptive field-section helpers (for CRD-XX sheets)
# ---------------------------------------------------------------------------

def find_field_section(ws):
    for row in ws.iter_rows(min_row=1, max_row=25, min_col=1, max_col=16):
        for cell in row:
            if cell.value and 'custom field detail' in str(cell.value).lower():
                h_row = cell.row
                h_col = cell.column
                for r in range(h_row + 1, h_row + 6):
                    v_same = str(ws.cell(row=r, column=h_col).value or '').strip()
                    v_next = str(ws.cell(row=r, column=h_col + 1).value or '').strip()
                    if v_same and ':' in v_same:
                        return h_row, h_col, h_col + 1
                    if v_next and ':' in v_next:
                        return h_row, h_col + 1, h_col + 2
                return h_row, h_col, h_col + 1
    return None


def extract_field_values(ws, header_row, label_col, value_col):
    label_map = {}
    for r in range(header_row + 1, header_row + 35):
        raw = str(ws.cell(row=r, column=label_col).value or '').strip()
        if not raw:
            continue
        normalized = raw.lower().rstrip(':').strip()
        label_map[normalized] = r

    def lookup(*terms):
        for term in terms:
            t = term.lower()
            for lbl, row in label_map.items():
                if lbl == t or lbl.startswith(t):
                    return ws.cell(row=row, column=value_col).value
        return None

    return {
        'document_type':        lookup('document'),
        'label':                lookup('label (key)', 'label'),
        'technical_name':       lookup('technical name'),
        'type':                 lookup('type'),
        'reportable':           lookup('reportable'),
        'erp_importable':       lookup('erp importable', 'erp import'),
        'erp_exportable':       lookup('erp exportable', 'erp export'),
        'an_field':             lookup('an field'),
        'visibility_condition': lookup('visibility condition', 'visibility'),
    }


# ---------------------------------------------------------------------------
# Format detection
# ---------------------------------------------------------------------------

NEW_FORMAT_SIGNAL_SHEETS = [
    'Sourcing_Customizations', 'Contracts_Customizations',
    'SPM_Customizations', 'Forms_Customizations', 'Savings_Customizations'
]

def detect_format(wb):
    """Returns 'new' if consolidated upstream sheets present, else 'old'."""
    for s in NEW_FORMAT_SIGNAL_SHEETS:
        if s in wb.sheetnames:
            return 'new'
    return 'old'


# ---------------------------------------------------------------------------
# Overview / status reader
# ---------------------------------------------------------------------------

def read_overview(wb, fmt):
    if fmt == 'new':
        ov_name = 'Overview' if 'Overview' in wb.sheetnames else None
    else:
        ov_name = next((n for n in ['CRD_Overview', 'CRD Overview'] if n in wb.sheetnames), None)

    overview_data = {}
    if not ov_name:
        return overview_data

    ws = wb[ov_name]
    data_start = 7
    for i, row in enumerate(ws.iter_rows(max_row=15, values_only=True), 1):
        if row[0] and 'crd' in str(row[0]).lower() and 'no' in str(row[0]).lower():
            data_start = i + 1
            break

    for row in ws.iter_rows(min_row=data_start, max_row=300, values_only=True):
        if not row[0]:
            continue
        identifier = str(row[0]).strip()
        if not identifier:
            continue
        overview_data[identifier] = {
            'title':        row[1] if len(row) > 1 else None,
            'solution':     row[2] if len(row) > 2 else None,
            'status':       row[6] if len(row) > 6 else None,
            'fields_count': row[7] if len(row) > 7 else None,
        }
    return overview_data


def get_status_for(crd_id, overview_data):
    if crd_id in overview_data:
        return overview_data[crd_id].get('status')
    crd_clean = re.sub(r'[\s\-_]', '', crd_id).lower()
    for key, val in overview_data.items():
        key_clean = re.sub(r'[\s\-_]', '', key).lower()
        if crd_clean == key_clean or key_clean.endswith(crd_clean):
            return val.get('status')
    return None


# ---------------------------------------------------------------------------
# B&I CRD-XX sheet extraction
# ---------------------------------------------------------------------------

def extract_bi_crds(wb, target_crds, overview_data):
    results = []
    for crd_name in target_crds:
        if crd_name not in wb.sheetnames:
            continue
        ws = wb[crd_name]

        crd_no         = ws['C8'].value
        title          = ws['C10'].value
        existing_field = ws['C13'].value

        section = find_field_section(ws)
        if section:
            h_row, lbl_col, val_col = section
            fv = extract_field_values(ws, h_row, lbl_col, val_col)
            doc_type   = fv['document_type']
            label      = fv['label']
            tech_name  = fv['technical_name']
            field_type = fv['type']
            reportable = fv['reportable']
            erp_import = fv['erp_importable']
            erp_export = fv['erp_exportable']
            an_field   = fv['an_field']
            vis_inline = fv['visibility_condition']
            layout_used = f'adaptive (label col {lbl_col}, value col {val_col})'
        else:
            doc_type   = ws['K10'].value
            label      = ws['K11'].value
            tech_name  = ws['K12'].value
            field_type = ws['K13'].value
            reportable = ws['K14'].value
            erp_import = ws['K15'].value
            erp_export = ws['K16'].value
            an_field   = ws['K17'].value
            vis_inline = None
            layout_used = 'fallback (K col)'

        visibility_condition = vis_inline or ws['J31'].value
        fmd_name       = ws['M10'].value
        relation_entry = ws['M20'].value

        desc_parts = []
        for r in range(15, 40):
            val = ws.cell(row=r, column=2).value
            if val and str(val).strip():
                desc_parts.append(str(val).strip())
        description = ' '.join(desc_parts)

        struck_off = False
        try:
            if (ws['C8'].font and ws['C8'].font.strikethrough) or \
               (ws['C10'].font and ws['C10'].font.strikethrough):
                struck_off = True
        except Exception:
            pass

        ov        = overview_data.get(crd_name, {})
        status    = ov.get('status') or get_status_for(crd_name, overview_data)
        cancelled = bool(status and str(status).strip().lower() in ['cancelled', 'rejected'])

        if existing_field and 'please answer' in str(existing_field).lower():
            existing_field = None

        has_field_data = bool(label or doc_type or field_type)
        fields = []
        if has_field_data:
            def clean(v):
                return str(v).strip() if v is not None else None
            fields.append({
                'document_type':    clean(doc_type),
                'label':            clean(label),
                'technical_name':   clean(tech_name),
                'type':             clean(field_type),
                'reportable':       clean(reportable),
                'erp_importable':   clean(erp_import),
                'erp_exportable':   clean(erp_export),
                'an_field':         clean(an_field),
                '_layout_detected': layout_used,
            })

        results.append({
            'crd_no':             crd_no or crd_name,
            'title':              title,
            'description':        description,
            'struck_off':         struck_off,
            'cancelled':          cancelled,
            'status':             status,
            'existing_field':     existing_field,
            'areas_impacted':     ov.get('solution', ''),
            'fields':             fields,
            'visibility_condition': str(visibility_condition).strip() if visibility_condition else None,
            'fmd':                {'name': str(fmd_name).strip()} if fmd_name else None,
            'relation_entry':     str(relation_entry).strip() if relation_entry else None,
        })
    return results


# ---------------------------------------------------------------------------
# Consolidated upstream sheet config
# ---------------------------------------------------------------------------

# id_pattern: only rows whose id_col value matches this pattern are kept
CONSOLIDATED_CFG = {
    'Sourcing_Customizations': {
        'scope': 'Upstream', 'document_type': 'Sourcing Header',
        'id_pattern': r'^S-\d+$',
        'header_row': 4, 'data_start': 6,
        'id_col': 2, 'title_col': 3, 'modification_col': 4,
        'label_col': 6, 'type_col': 7,
        'reportable_col': 9, 'required_col': 11,
        'condition_col': 14,
    },
    'Contracts_Customizations': {
        'scope': 'Upstream', 'document_type': 'Contract Header',
        'id_pattern': r'^C-\d+$',
        'header_row': 4, 'data_start': 6,
        'id_col': 2, 'title_col': 3, 'modification_col': 4,
        'label_col': 6, 'type_col': 7,
        'reportable_col': 9, 'required_col': 11,
        'condition_col': 14,
    },
    'SPM_Customizations': {
        'scope': 'Upstream', 'document_type': 'SPM Header',
        'id_pattern': r'^SPM-\d+$',
        'header_row': 4, 'data_start': 6,
        'id_col': 2, 'title_col': 3, 'modification_col': 4,
        'label_col': 6, 'type_col': 7,
        'reportable_col': 9, 'required_col': 11,
        'condition_col': 14,
    },
    'Forms_Customizations': {
        'scope': 'Upstream', 'document_type': 'dForm / Savings Form',
        'id_pattern': r'^F-\d+$',
        'header_row': 5, 'data_start': 7,
        'id_col': 2, 'title_col': 3, 'modification_col': 5,
        'label_col': 7, 'type_col': 9,
        'reportable_col': None, 'required_col': 11,
        'condition_col': 13,
    },
    'Savings_Customizations': {
        'scope': 'Upstream', 'document_type': 'Savings Form',
        'id_pattern': r'^SF-\d+$',
        'header_row': 2, 'data_start': 4,
        'id_col': 2, 'title_col': 3, 'modification_col': 4,
        'label_col': 7, 'type_col': 8,
        'reportable_col': None, 'required_col': 10,
        'condition_col': 11,
    },
}

PLACEHOLDER_TITLES = {
    'enter a short description', 'enter a title', '9999999', 'sample',
    'enter a title or this request. ', 'new sourcing field 1', 'new sourcing field 2',
    'new contract field 1', 'new contract field 2', 'new spm field 1', 'new spm field 2',
    'new forms field 1', 'new forms field 2', 'new savings form field 1', 'new savings form field 2',
    'new savings form field 3', 'new savings form field 4', 'new savings form field 5',
}

def is_placeholder_title(v):
    if v is None:
        return True
    s = str(v).strip()
    return s == '0' or s == '' or s.lower()[:50] in PLACEHOLDER_TITLES


def extract_consolidated_sheet(wb, sheet_name, cfg, overview_data, target_ids=None):
    """Extract customization rows from a consolidated upstream sheet."""
    if sheet_name not in wb.sheetnames:
        return []

    ws = wb[sheet_name]
    id_pat = re.compile(cfg['id_pattern'], re.IGNORECASE)
    results = []
    seen_ids = set()  # prevent duplicates

    def cv(r, col):
        if col is None:
            return None
        v = ws.cell(row=r, column=col).value
        return str(v).strip() if v is not None else None

    for r in range(cfg['data_start'], ws.max_row + 1):
        raw_id = cv(r, cfg['id_col'])
        if not raw_id:
            continue

        # Only accept IDs matching the expected pattern (filters junk rows)
        if not id_pat.match(raw_id):
            continue

        # Skip duplicates (translation section re-uses same IDs)
        if raw_id in seen_ids:
            continue
        seen_ids.add(raw_id)

        # Target ID filter
        if target_ids and raw_id not in target_ids:
            continue

        title        = cv(r, cfg['title_col'])
        label        = cv(r, cfg['label_col'])
        field_type   = cv(r, cfg['type_col'])
        reportable   = cv(r, cfg['reportable_col']) if cfg.get('reportable_col') else None
        condition    = cv(r, cfg['condition_col'])
        modification = cv(r, cfg['modification_col'])
        required     = cv(r, cfg['required_col'])

        # Use ID as fallback title if placeholder
        if is_placeholder_title(title):
            title = raw_id

        status    = get_status_for(raw_id, overview_data)
        cancelled = bool(status and str(status).strip().lower() in
                         ['cancelled but built', 'cancelled', 'rejected'])

        desc_parts = []
        if modification:
            desc_parts.append(f'Type: {modification}')
        if required:
            desc_parts.append(f'Required: {required}')
        if condition:
            desc_parts.append(f'Condition: {condition}')
        description = ' | '.join(desc_parts) if desc_parts else f'{sheet_name} customization'

        fields = []
        if label or field_type:
            fields.append({
                'document_type':    cfg['document_type'],
                'label':            label,
                'technical_name':   None,
                'type':             field_type,
                'reportable':       reportable,
                'erp_importable':   None,
                'erp_exportable':   None,
                'an_field':         None,
                '_layout_detected': f'consolidated-{sheet_name}',
            })

        results.append({
            'crd_no':             raw_id,
            'title':              title,
            'description':        description,
            'struck_off':         False,
            'cancelled':          cancelled,
            'status':             status,
            'existing_field':     'No' if modification and 'new' in str(modification).lower() else None,
            'areas_impacted':     cfg['scope'],
            'fields':             fields,
            'visibility_condition': condition,
            'fmd':                None,
            'relation_entry':     None,
        })

    return results


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

file_path = sys.argv[1]
crd_arg   = sys.argv[2].strip() if len(sys.argv) > 2 else 'all'

wb  = openpyxl.load_workbook(file_path, data_only=True)
fmt = detect_format(wb)
overview_data = read_overview(wb, fmt)

all_bi_sheets = [s for s in wb.sheetnames if re.match(r'^CRD-\d+$', s)]

results = []

if crd_arg.lower() == 'all':
    results += extract_bi_crds(wb, all_bi_sheets, overview_data)
    if fmt == 'new':
        for sheet_name, cfg in CONSOLIDATED_CFG.items():
            results += extract_consolidated_sheet(wb, sheet_name, cfg, overview_data)
else:
    requested = [c.strip() for c in crd_arg.split(',')]
    bi_req  = [r for r in requested if re.match(r'^CRD-\d+$', r, re.IGNORECASE)]
    ups_req = [r for r in requested if not re.match(r'^CRD-\d+$', r, re.IGNORECASE)]

    results += extract_bi_crds(wb, bi_req, overview_data)

    if fmt == 'new' and ups_req:
        for sheet_name, cfg in CONSOLIDATED_CFG.items():
            results += extract_consolidated_sheet(wb, sheet_name, cfg, overview_data,
                                                  target_ids=set(ups_req))

print(json.dumps(results, indent=2, default=str))
